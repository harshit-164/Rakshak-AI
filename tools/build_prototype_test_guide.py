from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "artifacts/Safeguard_AI_Prototype_Review_and_Observation_Test_Guide.docx"

NAVY = "15324B"
TEAL = "147D82"
PALE_TEAL = "E8F4F3"
PALE_BLUE = "EAF0F5"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
INK = "17242E"
MUTED = "5D6A73"
WHITE = "FFFFFF"
GRID = "C9D4DC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    if keep_lines:
        node = OxmlElement("w:keepLines")
        p_pr.append(node)


def set_run(run, size=None, bold=None, color=None, italic=None, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, *, bold=False, color=INK, size=10.5, after=6, before=0,
             align=WD_ALIGN_PARAGRAPH.LEFT, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    if keep:
        set_repeat_keep(p, keep_next=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=10.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1, *, page_break_before=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(text)
    set_run(r, size={1: 16, 2: 13, 3: 11.5}[level], bold=True,
            color={1: NAVY, 2: TEAL, 3: NAVY}[level])
    set_repeat_keep(p, keep_next=True)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size="0")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_run(r, size=9, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(text)
    set_run(r2, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_run(r, size=font_size, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows):
        added_row = table.add_row()
        tr_pr = added_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        cells = added_row.cells
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, "F6F8FA")
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_lines(doc, labels, line_count=2):
    for label in labels:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label}: ")
        set_run(r, size=9.5, bold=True, color=NAVY)
        r2 = p.add_run("_" * 82)
        set_run(r2, size=9.5, color=GRID)
        for _ in range(line_count - 1):
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(3)
            r3 = p2.add_run("_" * 98)
            set_run(r3, size=9.5, color=GRID)


def page_break(doc):
    # Major sections flow naturally so a short continuation is not stranded on
    # a nearly empty page. Heading keep-with-next rules preserve scanability.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)


def hard_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


doc = Document()
doc.settings.odd_and_even_pages_header_footer = True
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.38)
section.footer_distance = Inches(0.38)

# Compact reference guide token map.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, before, after, color in (
    ("Heading 1", 16, 18, 10, NAVY),
    ("Heading 2", 13, 14, 7, TEAL),
    ("Heading 3", 11.5, 10, 5, NAVY),
):
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Bullet 2", "List Number"):
    st = styles[name]
    st.font.name = "Aptos"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25

# Running header and footer. Populate both variants so Word and LibreOffice
# render the same furniture on every page.
for header in (section.header, section.even_page_header):
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("SAFEGUARD AI  /  PROTOTYPE TEST KIT")
    set_run(hr, size=8.5, bold=True, color=MUTED)

for footer in (section.footer, section.even_page_footer):
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Internal usability research  |  5 September 2026  |  ")
    set_run(fr, size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)

# Cover: workshop_agenda pattern.
add_text(doc, "SIH26165  /  RESEARCH KIT", bold=True, color=TEAL, size=10, after=24)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Prototype Review &\nObservation Test Guide")
set_run(r, size=28, bold=True, color=NAVY, font="Aptos Display")
add_text(doc, "Safeguard AI - SIF review assistant", size=14, color=TEAL, after=8)
add_text(doc, "A moderator-ready plan for observing whether reviewers can submit, interpret, challenge, and trust an evidence-grounded safety assessment.", size=11.5, color=MUTED, after=22)

metric = doc.add_table(rows=1, cols=3)
set_table_geometry(metric, [3120, 3120, 3120], indent_dxa=120)
set_table_borders(metric, color=TEAL, size="6")
for i, (big, small) in enumerate((("45 min", "SESSION"), ("5-8", "TARGET USERS"), ("7", "CORE TASKS"))):
    cell = metric.cell(0, i)
    set_cell_shading(cell, PALE_TEAL)
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(1)
    rr = p1.add_run(big)
    set_run(rr, size=16, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    rr2 = p2.add_run(small)
    set_run(rr2, size=8, bold=True, color=TEAL)

add_heading(doc, "Research decision", 2)
add_callout(doc, "Use this test to decide", "Can a safety reviewer understand what the prototype is proposing, verify its evidence, identify uncertainty, and know how to record a correction without treating the output as a safety verdict?", PALE_BLUE, NAVY)
add_text(doc, "Prepared from the repository's PRD, test plan, annotation guide, current interface, and implementation status. Test only with synthetic or cleared public reports.", size=9.5, color=MUTED, italic=True, after=12)

hard_page_break(doc)
add_heading(doc, "Document map", 2)
add_table(doc, ["Section", "Use"], [
    ("1. Test brief", "Scope, participants, success signals, and current-build boundary"),
    ("2. Run of show", "Set-up checklist and moderator script"),
    ("3. Participant tasks", "Seven realistic tasks with expected observations"),
    ("4. Observation sheets", "Per-task scoring and note capture"),
    ("5. Debrief", "Trust, comprehension, and adoption questions"),
    ("6. Synthesis", "Issue severity, prioritization, and go/no-go decision"),
], [1800, 7560], font_size=9.2)

page_break(doc)
add_heading(doc, "1. Test brief", 1)
add_heading(doc, "Objective", 2)
add_text(doc, "Evaluate the end-to-end reviewer experience, not model accuracy alone. Watch whether participants can form a correct mental model, complete the working observation flow, inspect exact-text evidence, distinguish yes/no/unknown, and recognize that the machine output still requires human review.")

add_heading(doc, "Current-build boundary", 2)
add_callout(doc, "Known limitation - do not conceal", "The current web UI exposes observation submission, progress, saved-result loading, assessment details, and provenance. It labels the result 'review pending,' but does not yet expose controls to save a human correction. Dashboard, CSV import, export, and the complete review loop remain specified or backend-oriented work, not current visible tasks.", PALE_GOLD, "7A5A00")

add_heading(doc, "Participants", 2, page_break_before=True)
add_bullet(doc, "Primary: 5-8 HSE reviewers, safety officers, supervisors, or people who routinely triage incident/near-miss narratives.")
add_bullet(doc, "Secondary: 2-3 non-HSE stakeholders to test whether risk language and limitations are understandable outside the core domain.")
add_bullet(doc, "Avoid using the project team as the only sample; familiarity hides navigation and terminology problems.")

add_heading(doc, "Research questions", 2)
for q in (
    "Do users understand that the SIF label is a proposal, not a verdict or work authorization?",
    "Can users submit a useful narrative and understand which optional fields affect the model?",
    "Can they trace the assessment back to exact submitted evidence?",
    "Can they distinguish an honest unknown from no and from a system/provider failure?",
    "Do rule tags, barrier states, provenance, and recorded/live status support appropriate trust?",
    "When review controls are absent, do users notice and correctly describe what they need to finish the job?",
): add_bullet(doc, q)

add_heading(doc, "Success signals", 2)
add_table(doc, ["Signal", "Target for pilot", "Failure to flag"], [
    ("Critical comprehension", "100% say output requires human review", "Any participant treats result as approval"),
    ("Core task completion", ">=80% unassisted on working tasks", "Repeated moderator rescue"),
    ("Evidence verification", ">=80% locate and explain exact-text evidence", "Evidence confused with general advice"),
    ("Uncertainty", ">=80% distinguish unknown, no, and error", "Unknown interpreted as safe/no-risk"),
    ("Review readiness", ">=80% describe a correction + reason", "Cannot find a way to finish review"),
    ("Ease", "Median single-ease score >=5/7", "Any core task median <=3/7"),
], [2340, 3240, 3780], font_size=8.6)

add_heading(doc, "2. Run of show", 1, page_break_before=True)
add_heading(doc, "Session setup", 2)
for item in (
    "[ ] Use a clean browser or isolated demo session.",
    "[ ] Confirm the tested build, URL, date, model revision, and whether inference is live or recorded.",
    "[ ] Prepare only synthetic scenarios from this guide; do not paste internal incident data.",
    "[ ] Start screen/audio recording only after consent; never record sensitive report content.",
    "[ ] Keep one observer silent and one moderator; capture timestamps for notable moments.",
    "[ ] Verify the app's configured/unconfigured state before the participant arrives.",
): add_bullet(doc, item)

add_heading(doc, "45-minute agenda", 2)
add_table(doc, ["Time", "Activity", "Moderator intent"], [
    ("0-4 min", "Welcome, consent, context", "Set expectations; no product teaching"),
    ("4-7 min", "Background questions", "Understand role and current triage process"),
    ("7-10 min", "First impression", "Test value proposition and boundaries"),
    ("10-30 min", "Tasks 2-6", "Think-aloud observation; minimal prompting"),
    ("30-35 min", "Review concept/gap task", "Learn expected correction workflow"),
    ("35-42 min", "Debrief", "Probe trust, clarity, usefulness, missing controls"),
    ("42-45 min", "Ease ratings + close", "Collect ratings before discussion drifts"),
], [1200, 2820, 5340], font_size=9)

add_heading(doc, "Opening script", 2)
add_callout(doc, "Read aloud", "We are testing the prototype, not you. Please say what you expect, notice, and find confusing. The system is experimental and does not certify that work is safe. Use only the synthetic examples I provide. I may stay quiet so I can see what the interface communicates on its own.", PALE_TEAL, TEAL)

add_heading(doc, "Neutral prompts", 2)
for prompt in (
    "What are you thinking now?",
    "What did you expect to happen?",
    "What does that label mean to you?",
    "Where would you look next?",
    "What, if anything, makes you confident in that conclusion?",
): add_bullet(doc, prompt)
add_text(doc, "Avoid: 'Click Analyze,' 'That quote proves it,' or any prompt that explains the intended answer before the participant responds.", bold=True, color="9B1C1C", size=9.5)

page_break(doc)
add_heading(doc, "3. Participant tasks", 1)
add_text(doc, "Give tasks one at a time. Do not show the expected behavior column to the participant. Record outcome before discussing quality.", italic=True, color=MUTED)

add_heading(doc, "Task 1 - First impression", 2)
add_callout(doc, "Participant prompt", "Without clicking yet, tell me what you think this product does, who it is for, and what it should not be used for.", PALE_BLUE, NAVY)
add_bullet(doc, "Observe: understanding of SIF review, human oversight, evidence grounding, and allowed data.")
add_bullet(doc, "Pass: participant describes prioritization/review support and does not describe certification or automatic safety approval.")

add_heading(doc, "Task 2 - Submit a serious-exposure observation", 2)
add_callout(doc, "Synthetic narrative", "During a lifting operation, a suspended steel bundle passed directly above a worker who had entered the barricaded area. The lift was stopped before contact and nobody was injured.", PALE_TEAL, TEAL)
add_text(doc, "Ask the participant to submit it with Activity = Mechanical lifting, Site = Demo Site Alpha, today's test date, and the synthetic-content box selected.")
add_bullet(doc, "Observe: field meaning, synthetic-data warning, optional-field comprehension, confidence before submit, and progress-state clarity.")
add_bullet(doc, "Expected concept: credible SIF potential; line of fire and safe mechanical lifting may be relevant. The exact model response may differ and must not be coached.")

add_heading(doc, "Task 3 - Inspect and challenge the result", 2)
add_callout(doc, "Participant prompt", "Decide whether you would prioritize this report for human review. Show me which parts of the page support your decision and which parts you would challenge.", PALE_BLUE, NAVY)
add_bullet(doc, "Observe: whether exact quotes are found, narrative and evidence are distinguished, rule tags are understood, and barriers are not over-interpreted.")
add_bullet(doc, "Ask after completion: What do 'model,' 'revision,' 'prompt,' 'rubric,' and 'live/recorded' tell you?")

add_heading(doc, "Task 4 - Test an honest unknown", 2)
add_callout(doc, "Synthetic narrative", "A lifting issue occurred at Demo Site Beta. Details about the load, worker position, and controls were unavailable.", PALE_TEAL, TEAL)
add_text(doc, "Ask the participant to start a new analysis and explain the result. Do not say that unknown is expected.")
add_bullet(doc, "Pass concept: participant recognizes missing exposure/control facts and does not treat unknown as no or safe.")

page_break(doc)
add_heading(doc, "Task 5 - Distinguish low potential from uncertainty", 2)
add_callout(doc, "Synthetic narrative", "An office employee received a shallow paper cut while sorting paper files at a desk. No powered equipment or other hazard was involved.", PALE_TEAL, TEAL)
add_text(doc, "Ask: How is the meaning of this result different from the previous result? What action, if any, would you take?")
add_bullet(doc, "Pass concept: no means the supplied narrative supports low potential under the pilot rubric; it is not permission to proceed with work.")

add_heading(doc, "Task 6 - Reopen and privacy boundary", 2)
add_callout(doc, "Participant prompt", "Refresh this report and check whether the saved assessment is still available. Then tell me what you expect would happen if you opened its URL in another anonymous browser session.", PALE_BLUE, NAVY)
add_bullet(doc, "Observe: persistence expectations, loading feedback, and session-privacy mental model.")
add_bullet(doc, "Expected: same session can reopen; a different session must not read the report and should receive a not-found/unauthorized-safe experience.")

add_heading(doc, "Task 7 - Complete a human review (concept/gap test)", 2)
add_callout(doc, "Participant prompt", "Assume you disagree with the SIF label or rule tags. Show me how you would correct the assessment, record your reason, and confirm the original model result remains available.", PALE_GOLD, "7A5A00")
add_text(doc, "Moderator instruction: Do not invent or point to controls. Let the participant search. When they conclude the action is unavailable, reveal that this is a known current-build gap and ask them to sketch or describe the minimum acceptable review flow.", size=9.7)
add_bullet(doc, "Capture expected fields: SIF label, relevant rules, assessed/unassessed state, confirmed breaches, barrier tags, evidence edits, reason/note, and save confirmation.")
add_bullet(doc, "Capture audit expectations: visible original prediction, review author/time/version, conflict warning, and behavior after a narrative revision.")

add_heading(doc, "Optional resilience probes", 2)
add_table(doc, ["Probe", "What to observe"], [
    ("Configuration unavailable", "Does the setup state explain what is blocked without exposing secrets?"),
    ("Provider failure", "Is failure distinct from a no/unknown result, and is safe retry behavior clear?"),
    ("Long/short narrative", "Are limits explained before or at the point of error?"),
    ("Mobile + keyboard", "Can all visible actions be reached, focused, read, and activated?"),
    ("Prompt injection text", "Does the model remain evidence-grounded and does UI avoid treating instructions as trusted?"),
], [2880, 6480], font_size=9)

page_break(doc)
add_heading(doc, "4. Observation sheet", 1)
add_text(doc, "Use one copy per participant. Record behavior and exact words before interpretation.", italic=True, color=MUTED)
add_table(doc, ["Session ID", "Role / experience", "Build + model", "Date / moderator"], [
    ("________________", "________________________", "________________________", "________________________"),
], [1500, 2760, 2580, 2520], font_size=8.7)

add_heading(doc, "Task scorecard", 2)
add_table(doc, ["Task", "Outcome", "Ease 1-7", "Time", "Critical observation / quote"], [
    ("1 First impression", "S / P / F", "__", "__", "________________________________"),
    ("2 Submit serious exposure", "S / P / F", "__", "__", "________________________________"),
    ("3 Inspect and challenge", "S / P / F", "__", "__", "________________________________"),
    ("4 Honest unknown", "S / P / F", "__", "__", "________________________________"),
    ("5 No vs unknown", "S / P / F", "__", "__", "________________________________"),
    ("6 Reopen/privacy", "S / P / F", "__", "__", "________________________________"),
    ("7 Human review gap", "N/A gap", "__", "__", "________________________________"),
], [1600, 1200, 950, 700, 4910], font_size=8.2)
add_text(doc, "Outcome codes: S = unassisted success; P = partial or prompted; F = failed/abandoned. Ease: 1 = very difficult, 7 = very easy. Score Task 7 for discoverability/expectation, not completion.", size=8.7, color=MUTED)

add_heading(doc, "Behavioral observations", 2)
add_lines(doc, ["First hesitation or wrong turn", "Evidence used to build trust", "Term or label misunderstood", "Unsafe interpretation, if any", "Workaround attempted"], line_count=1)

add_heading(doc, "Comprehension checks", 2)
for item in (
    "[ ] States that machine result requires human review",
    "[ ] Distinguishes unknown from no",
    "[ ] Distinguishes provider failure from a safety label",
    "[ ] Understands exact-text evidence",
    "[ ] Notices live vs recorded provenance",
    "[ ] Expects original model output to remain after correction",
): add_bullet(doc, item)

page_break(doc)
add_heading(doc, "5. Debrief", 1)
add_heading(doc, "Interview questions", 2)
questions = [
    "In your own words, what decision does this prototype help you make?",
    "What is the most useful part of the result page? What is the least useful?",
    "What would make you distrust or stop using an assessment?",
    "Which evidence would you need before changing a model label?",
    "What should happen when information is insufficient?",
    "What should the review screen capture, and what must remain immutable?",
    "Where would this fit into your current workflow, if anywhere?",
    "What data would you refuse to enter into this demo?",
    "What is the single most important improvement before another test round?",
]
for q in questions: add_number(doc, q)

add_heading(doc, "Trust calibration ratings", 2)
add_text(doc, "Circle one number for each statement: 1 = strongly disagree; 7 = strongly agree.", size=9.2, color=MUTED)
add_table(doc, ["Statement", "Rating"], [
    ("I can tell what the system knows versus what it is missing.", "1  2  3  4  5  6  7"),
    ("I can verify why the system proposed its SIF label.", "1  2  3  4  5  6  7"),
    ("I understand that the result is not a safety approval.", "1  2  3  4  5  6  7"),
    ("I know how I would correct the result and document why.", "1  2  3  4  5  6  7"),
    ("The provenance shown is meaningful to me.", "1  2  3  4  5  6  7"),
    ("I would use this to prioritize a review queue in a bounded pilot.", "1  2  3  4  5  6  7"),
], [6840, 2520], font_size=9)

add_heading(doc, "Participant's closing statement", 2)
add_lines(doc, ["What I would change first"], line_count=3)

page_break(doc)
add_heading(doc, "6. Synthesis and decision", 1)
add_heading(doc, "Issue log", 2)
add_table(doc, ["ID", "Evidence", "Impact", "Severity", "Owner / decision"], [
    ("__", "Observed behavior + quote", "Task/user/risk", "0 / 1 / 2 / 3", "________________"),
    ("__", "______________________", "____________", "0 / 1 / 2 / 3", "________________"),
    ("__", "______________________", "____________", "0 / 1 / 2 / 3", "________________"),
    ("__", "______________________", "____________", "0 / 1 / 2 / 3", "________________"),
    ("__", "______________________", "____________", "0 / 1 / 2 / 3", "________________"),
], [540, 3300, 1740, 1080, 2700], font_size=8.2)

add_heading(doc, "Severity rubric", 2)
add_table(doc, ["Level", "Definition", "Examples"], [
    ("0 - Note", "Preference or observation with no clear task impact", "Copy preference; optional enhancement"),
    ("1 - Minor", "Slows a user but recovery is easy", "Small label ambiguity; extra navigation"),
    ("2 - Major", "Blocks a core task or creates material misunderstanding", "Cannot verify evidence; cannot recover"),
    ("3 - Critical", "Could enable unsafe interpretation, privacy breach, or false product claim", "Unknown read as safe; cross-session access; recorded shown as live"),
], [1080, 3780, 4500], font_size=8.6)

add_heading(doc, "Release decision rules", 2, page_break_before=True)
add_callout(doc, "Stop / fix before broader demo", "Any Severity 3 finding; any cross-session data exposure; any participant interprets the output as authorization and the interface fails to correct them; or live/recorded/model identity is misleading.", PALE_RED, "9B1C1C")
add_callout(doc, "Iterate and retest", "Two or more participants fail the same core working task, median ease is 3/7 or lower, or reviewers cannot state the minimum correction workflow.", PALE_GOLD, "7A5A00")
add_callout(doc, "Ready for next bounded round", "No critical findings; at least 80% unassisted completion on working core tasks; all participants understand human-review and data-use limits; known review UI gap has an agreed design and implementation owner.", PALE_TEAL, TEAL)

add_heading(doc, "Session summary", 2)
add_table(doc, ["Measure", "Result", "Decision / next action"], [
    ("Participants / roles", "________________", "____________________________"),
    ("Core completion", "____ / ____", "____________________________"),
    ("Median ease", "____ / 7", "____________________________"),
    ("Critical / major issues", "____ / ____", "____________________________"),
    ("Review UI decision", "________________", "____________________________"),
    ("Retest owner + date", "________________", "____________________________"),
], [2700, 2100, 4560], font_size=8.3)

add_text(doc, "Basis: repository PRD, test plan, annotation guide, status, and current Analyze/Report screens. Usability test only - not an accuracy or operational-safety validation.", size=7.5, color=MUTED, italic=True, before=3, after=0)

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
